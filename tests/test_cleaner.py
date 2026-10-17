import io
import json
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from data_cleaner import web
from data_cleaner.cleaner import Options, clean, profile, to_xlsx
from data_cleaner.cli import main
from data_cleaner.detect import detect_date_order, detect_decimal, detect_kind
from data_cleaner.table import Table, TableError, read_any, read_csv, write_csv

SAMPLES = Path(__file__).resolve().parents[1] / "samples"


def table(columns, *rows):
    return Table(list(columns), [list(r) for r in rows])


# --- reading --------------------------------------------------------------------------------


def test_csv_encoding_and_delimiter_are_detected():
    text = "Nombre;Ciudad\r\nJosé;Medellín\r\n;\r\nAna;Bogotá\r\n"
    for data in (text.encode("utf-8-sig"), text.encode("cp1252")):
        t = read_csv(data)
        assert t.columns == ["Nombre", "Ciudad"]
        assert t.rows == [["José", "Medellín"], ["Ana", "Bogotá"]]
        assert t.source_rows == [2, 4]  # the blank line keeps its spreadsheet number


def test_ragged_rows_and_duplicate_headers_are_squared():
    t = read_csv(b"a,a,\n1,2\n3,4,5,6\n")
    assert t.columns == ["a", "a_2", "column_3", "column_4"]
    assert t.rows[0] == ["1", "2", "", ""]


def test_xlsx_dates_and_numbers_become_text():
    from datetime import datetime

    wb = Workbook()
    wb.active.append(["Fecha", "Qty"])
    wb.active.append([datetime(2026, 3, 5), 3.0])
    buf = io.BytesIO()
    wb.save(buf)
    t = read_any(buf.getvalue(), "x.xlsx")
    assert t.rows == [["2026-03-05", "3"]]


def test_unsupported_file_type():
    with pytest.raises(TableError, match="Unsupported"):
        read_any(b"{}", "data.json")


# --- detection ------------------------------------------------------------------------------


def test_kind_detection_from_headers_and_values():
    assert detect_kind("Correo", ["a@b.co", "c@d.com"]) == "email"
    assert detect_kind("Teléfono", ["300 123 4567", "+57 3001234567"]) == "phone"
    assert detect_kind("Fecha de registro", ["25/03/2024", "2024-04-01"]) == "date"
    assert detect_kind("Order ID", ["1001", "1002"]) == "id"
    assert detect_kind("Valor compra", ["$ 1.250.000", "320.000"]) == "number"
    assert detect_kind("Customer Name", ["ANA PEREZ"]) == "name"
    assert detect_kind("contact", ["x@y.com"] * 5) == "email"  # values beat a wrong hint
    assert detect_kind("Fecha", ["mañana", "ayer", "hoy"]) == "text"  # hint not confirmed
    assert detect_kind("misc", ["Activo", "activo", "Inactivo"] * 5) == "category"
    assert detect_kind("notes", ["called twice", "left a message", "wants a demo"]) == "text"


def test_date_order_and_decimal_detection():
    assert detect_date_order(["03/04/2026", "25/12/2025"], "MDY") == "DMY"
    assert detect_date_order(["03/04/2026", "12/25/2025"], "DMY") == "MDY"
    assert detect_date_order(["03/04/2026"], "MDY") == "MDY"
    assert detect_decimal(["$ 1.250.000", "320.000,50"], ".") == ","
    assert detect_decimal(["$1,250.00", "99.90"], ",") == "."


# --- cleaning -------------------------------------------------------------------------------

CONTACTS = table(
    ["Nombre", "Correo", "Teléfono", "Ciudad", "Fecha", "Valor"],
    [" ANA  PEREZ ", "Ana@Gmial.com", "300 123 4567", "medellin", "25/03/2024", "$ 1.250.000"],
    ["Luis Gómez", "luis@outlook.com", "N/A", "Medellín", "03/04/2024", "320.000,50"],
    ["ana perez", "ANA@GMAIL.COM", "+57 300 123 4567", "Bogotá", "25/03/2024", "1250000"],
    ["Sofía Ríos", "sofia at mail.com", "12345", "Medellín", "31/02/2024", "TBD"],
    ["Luis Gómez", "luis@outlook.com", "-", "Medellín", "03/04/2024", "320.000,50"],
)


def test_clean_end_to_end():
    out, report = clean(CONTACTS, Options(country="CO"))
    assert out.rows[0] == [
        "Ana Perez",
        "ana@gmail.com",
        "+573001234567",
        "Medellín",
        "2024-03-25",
        "1250000",
    ]
    assert out.rows[1][2] == ""  # "N/A" became empty
    assert out.rows[1][5] == "320000.5"

    # Row 6 repeats row 3 exactly once cleaned, so it is removed and reported.
    assert report.rows_in == 5 and report.rows_out == 4
    assert report.duplicates == [{"row": 6, "duplicate_of": 3}]

    problems = {(i.row, i.column): i.problem for i in report.issues}
    assert problems[(5, "Correo")] == "not a valid email address"
    assert problems[(5, "Teléfono")] == "does not look like a CO phone number"
    assert problems[(5, "Fecha")] == "unrecognized or impossible date"
    assert problems[(5, "Valor")] == "not a number"
    assert problems[(4, "Correo")] == "same email as row 2: possible duplicate"

    change = next(c for c in report.changes if c.row == 2 and c.column == "Correo")
    assert (change.before, change.after) == ("Ana@Gmial.com", "ana@gmail.com")
    assert change.rules == ["email_format", "email_domain_typo"]
    assert report.rule_counts["duplicate_row"] == 1


def test_options_change_the_outcome():
    _, flagged = clean(CONTACTS, Options(duplicates="flag", flag_shared_contacts=False))
    assert flagged.rows_out == 5
    assert [i.problem for i in flagged.issues if i.row == 6] == ["duplicate of row 3"]
    assert not any("possible duplicate" in i.problem for i in flagged.issues)

    out, _ = clean(CONTACTS, Options(kinds={"Valor": "text", "Nombre": "text"}))
    assert out.rows[0][0] == "ANA PEREZ"  # only whitespace is touched on text columns
    assert out.rows[0][5] == "$ 1.250.000"

    out, _ = clean(CONTACTS, Options(date_order="MDY"))
    assert out.rows[1][4] == "2024-03-04"


def test_category_variants_merge_to_the_best_spelling():
    t = table(["Ciudad"], *[[v] for v in ["bogota", "BOGOTÁ", "Bogota", "Bogota", "Bogotá"] * 3])
    out, _ = clean(t)
    assert {r[0] for r in out.rows} == {"Bogotá"}


def test_profile_reports_examples():
    cols = profile(CONTACTS)
    assert [c.kind for c in cols] == ["name", "email", "phone", "category", "date", "number"]
    assert cols[0].examples[0] == "ANA PEREZ"


# --- output ---------------------------------------------------------------------------------


def test_xlsx_output_has_typed_values_and_audit_sheets():
    out, report = clean(CONTACTS)
    wb = load_workbook(io.BytesIO(to_xlsx(out, report)))
    assert wb.sheetnames == ["Cleaned", "Issues", "Changes"]
    first = [c.value for c in wb["Cleaned"][2]]
    assert first[4].year == 2024 and first[5] == 1250000
    assert wb["Issues"].max_row == len(report.issues) + 1


def test_csv_output_has_bom_for_excel():
    out, _ = clean(CONTACTS)
    assert write_csv(out).startswith(b"\xef\xbb\xbfNombre,")


# --- samples, CLI and browser bridge -------------------------------------------------------


def test_sample_contacts():
    t = read_any((SAMPLES / "crm-contacts.csv").read_bytes(), "crm-contacts.csv")
    out, report = clean(t)
    assert report.rows_in == 247 and len(report.duplicates) == 13
    cities = {r[3] for r in out.rows}
    assert cities == {"Medellín", "Bogotá", "Cali", "Barranquilla", "Envigado"}


def test_cli_clean(tmp_path, capsys):
    out, rep = tmp_path / "out.xlsx", tmp_path / "report.json"
    code = main(
        [
            "clean",
            str(SAMPLES / "sales-orders.xlsx"),
            "-o",
            str(out),
            "--country",
            "US",
            "--allow-repeats",
            "--report",
            str(rep),
        ]
    )
    assert code == 0 and out.exists()
    report = json.loads(rep.read_text())
    assert {i["value"] for i in report["issues"]} == {"TBD", "13/31/2026"}
    assert "146 rows in, 140 rows out" in capsys.readouterr().err


def test_web_bridge_round_trip():
    data = (SAMPLES / "sales-orders.xlsx").read_bytes()
    loaded = json.loads(web.load(data, "sales-orders.xlsx"))
    assert loaded["rows"] == 146 and loaded["sheets"] == ["Orders"]
    result = json.loads(
        web.run(json.dumps({"country": "US", "flag_shared_contacts": False, "unknown_option": 1}))
    )
    assert result["rows_out"] == 140 and len(result["preview"]) == 140
    assert web.export("csv").startswith(b"\xef\xbb\xbf")
    assert json.loads(web.export("report"))["rows_in"] == 146
