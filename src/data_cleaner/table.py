"""Reading and writing tables. Everything is read as text so no rule sees a pre-mangled value.

CSV: encoding (UTF-8, UTF-8 with BOM, Windows-1252) and delimiter (, ; tab |) are detected.
XLSX: the first sheet (or a named one); real Excel dates become YYYY-MM-DD text.
"""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field
from datetime import date, datetime

ENCODINGS = ("utf-8-sig", "cp1252")
DELIMITERS = ",;\t|"


@dataclass
class Table:
    columns: list[str]
    rows: list[list[str]]
    source_rows: list[int] = field(default_factory=list)  # spreadsheet row of each row

    def __post_init__(self):
        if not self.source_rows:
            self.source_rows = list(range(2, len(self.rows) + 2))

    def column(self, i: int) -> list[str]:
        return [r[i] for r in self.rows]


class TableError(Exception):
    pass


def _decode(data: bytes) -> str:
    for enc in ENCODINGS:
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    raise TableError("Could not decode the file as UTF-8 or Windows-1252 text.")


def _square(header: list[str], body: list[list[str]]) -> Table:
    width = max([len(header), *(len(r) for r in body)]) if body else len(header)
    columns = [
        h.strip() or f"column_{i + 1}" for i, h in enumerate(header + [""] * (width - len(header)))
    ]
    seen: dict[str, int] = {}
    for i, name in enumerate(columns):  # duplicate headers would collide in dict-based tools
        if name in seen:
            seen[name] += 1
            columns[i] = f"{name}_{seen[name]}"
        else:
            seen[name] = 1
    rows, source = [], []
    for n, r in enumerate(body, start=2):
        if any(c.strip() for c in r):  # fully blank rows are dropped, keeping row numbers
            rows.append(r + [""] * (width - len(r)))
            source.append(n)
    return Table(columns, rows, source)


def read_csv(data: bytes) -> Table:
    text = _decode(data)
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=DELIMITERS)
        delimiter = dialect.delimiter
    except csv.Error:
        delimiter = max(DELIMITERS, key=lambda d: sample.split("\n", 1)[0].count(d))
    records = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    if not records:
        raise TableError("The file is empty.")
    return _square(records[0], records[1:])


def _cell(value) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return (
            value.date().isoformat()
            if value.time() == datetime.min.time()
            else value.isoformat(sep=" ")
        )
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def sheet_names(data: bytes) -> list[str]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True)
    names = wb.sheetnames
    wb.close()
    return names


def read_xlsx(data: bytes, sheet: str | None = None) -> Table:
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated types for bad files
        raise TableError(f"Could not open the Excel file: {exc}") from exc
    ws = wb[sheet] if sheet else wb.worksheets[0]
    records = [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    wb.close()
    if not records:
        raise TableError("The sheet is empty.")
    return _square(records[0], records[1:])


def read_any(data: bytes, filename: str, sheet: str | None = None) -> Table:
    name = filename.lower()
    if name.endswith((".xlsx", ".xlsm")):
        return read_xlsx(data, sheet)
    if name.endswith((".csv", ".txt", ".tsv")):
        return read_csv(data)
    raise TableError("Unsupported file type. Use .csv, .tsv, .txt or .xlsx.")


def write_csv(table: Table) -> bytes:
    """UTF-8 with BOM, so Excel opens accents correctly."""
    out = io.StringIO()
    writer = csv.writer(out, lineterminator="\n")
    writer.writerow(table.columns)
    writer.writerows(table.rows)
    return out.getvalue().encode("utf-8-sig")


def write_xlsx(sheets: dict[str, tuple[list[str], list[list]]]) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for title, (columns, rows) in sheets.items():
        ws = wb.create_sheet(title)
        ws.append(columns)
        for c in ws[1]:
            c.font = Font(bold=True)
        for row in rows:
            ws.append(row)
        ws.freeze_panes = "A2"
        for i, name in enumerate(columns, start=1):
            longest = max([len(str(name)), *(len(str(r[i - 1])) for r in rows[:500])])
            ws.column_dimensions[ws.cell(1, i).column_letter].width = min(48, longest + 2)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()
